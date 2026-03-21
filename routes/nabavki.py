# routes/nabavki.py
import sqlite3
import json
import os
import time
from datetime import datetime, date
from io import BytesIO
from flask import Blueprint, render_template, request, flash, redirect, url_for, jsonify, send_file, session, current_app
from PIL import Image as PILImage
from utils.db import get_db
from utils.decorators import login_required
from utils.helpers import _format_cet
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from utils.notifications import send_push_to_user, send_push_to_nabavki_group
from utils.nabavki_email import (
    notify_novo_baranje,
    notify_promena_status,
    notify_nov_komentar,
)

nabavki_bp = Blueprint('nabavki', __name__, url_prefix='/nabavki')

_IMG_MAX_SIZE = (1200, 1200)
_IMG_QUALITY  = 72


# ─────────────────────────────────────────────────────────────
# ПОМОШНА — зачувај и компресирај слика
# ─────────────────────────────────────────────────────────────
def _save_compressed_image(file_storage, save_dir, filename_base):
    try:
        os.makedirs(save_dir, exist_ok=True)
        final_name = f"{filename_base}.jpg"
        save_path  = os.path.join(save_dir, final_name)
        img = PILImage.open(file_storage)
        if img.mode in ("RGBA", "P", "LA"):
            img = img.convert("RGB")
        img.thumbnail(_IMG_MAX_SIZE, PILImage.Resampling.LANCZOS)
        img.save(save_path, format="JPEG", quality=_IMG_QUALITY, optimize=True)
        print(f"[SLIKA] Зачувана: {save_path}")
        return final_name
    except Exception as e:
        print(f"[SLIKA] Грешка при зачувување: {e}")
        return None


# ─────────────────────────────────────────────────────────────
# ПОМОШНА — осигури колона slika во коментари
# ─────────────────────────────────────────────────────────────
def _ensure_comment_slika_column(cursor):
    try:
        cursor.execute("ALTER TABLE nabavki_comments ADD COLUMN slika TEXT DEFAULT NULL")
    except Exception:
        pass


# ─────────────────────────────────────────────────────────────
# ПОМОШНА — осигури табела за архивирани коментари
# ─────────────────────────────────────────────────────────────
def _ensure_archive_comments_table(cursor):
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS nabavki_archive_comments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            archive_req_id INTEGER NOT NULL,
            user TEXT,
            comment TEXT,
            slika TEXT,
            timestamp TEXT
        )
    """)


# ─────────────────────────────────────────────────────────────
# NABAVKI MAIN PAGE
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/", methods=["GET", "POST"])
@login_required
def nabavki():
    conn            = get_db()
    cursor          = conn.cursor()
    _ensure_comment_slika_column(cursor)
    is_nabavki_user = session.get("is_admin") or session.get("user_group") == "Nabavki"

    if request.method == "POST":
        action = request.form.get("action")
        if action == "kreiraj":
            naslov       = request.form.get("naslov", "").strip()
            kolicina_str = request.form.get("kolicina", "0").strip()
            datum_itnost = request.form.get("datum_itnost")
            chat_comment = request.form.get("chat_comment", "").strip()

            slika_file      = request.files.get("slika")
            camera_filename = request.form.get("camera_slika_filename", "").strip()

            try:
                kolicina = int(kolicina_str)
                if not naslov or kolicina < 1 or not datum_itnost:
                    flash("Наслов, количина и датум на итност се задолжителни!", "error")
                    return redirect(url_for("nabavki.nabavki"))

                save_dir       = os.path.join(current_app.config['STATIC_FOLDER'], "nabavki")
                slika_filename = None
                filename_base  = f"req_{int(time.time())}_{session['user']}"

                if slika_file and slika_file.filename:
                    slika_filename = _save_compressed_image(slika_file, save_dir, filename_base)
                    if not slika_filename:
                        flash("Грешка при зачувување на сликата!", "warning")

                elif camera_filename:
                    from utils.config import POZICII_FOLDER
                    src = os.path.join(POZICII_FOLDER, camera_filename)
                    if os.path.exists(src):
                        dst_name = f"{filename_base}_cam.jpg"
                        dst      = os.path.join(save_dir, dst_name)
                        os.makedirs(save_dir, exist_ok=True)
                        try:
                            img = PILImage.open(src)
                            if img.mode in ("RGBA", "P", "LA"):
                                img = img.convert("RGB")
                            img.thumbnail(_IMG_MAX_SIZE, PILImage.Resampling.LANCZOS)
                            img.save(dst, format="JPEG", quality=_IMG_QUALITY, optimize=True)
                            slika_filename = dst_name
                        except Exception as e:
                            print(f"[SLIKA] Грешка при камера слика: {e}")
                    else:
                        flash(f"Камера сликата не е пронајдена: {camera_filename}", "warning")

                cursor.execute("""
                    INSERT INTO nabavki_requests
                    (username, naslov, kolicina, datum_kreiranje, datum_itnost, opis, slika, status)
                    VALUES (?, ?, ?, CURRENT_TIMESTAMP, ?, '', ?, 'креирано')
                """, (session["user"], naslov, kolicina, datum_itnost, slika_filename))
                conn.commit()
                new_id = cursor.lastrowid

                cursor.execute("""
                    CREATE TABLE IF NOT EXISTS fer_sequence (
                        id INTEGER PRIMARY KEY, last_num INTEGER DEFAULT 0
                    )
                """)
                cursor.execute("INSERT OR IGNORE INTO fer_sequence (id, last_num) VALUES (1, 0)")
                cursor.execute("UPDATE fer_sequence SET last_num = last_num + 1 WHERE id = 1")
                seq_row = cursor.execute("SELECT last_num FROM fer_sequence WHERE id = 1").fetchone()
                nalog   = f"Fer{seq_row['last_num']:03d}"
                cursor.execute("UPDATE nabavki_requests SET nalog_broj=? WHERE id=?", (nalog, new_id))
                conn.commit()

                if chat_comment:
                    cursor.execute("""
                        INSERT INTO nabavki_comments (req_id, user, comment) VALUES (?, ?, ?)
                    """, (new_id, session["user"], chat_comment))
                    conn.commit()

                now_local = datetime.now().strftime("%d-%m-%Y %H:%M")
                flash(f"Барањето е креирано со број: <strong>{nalog}</strong> на {now_local}!", "success")

                send_push_to_nabavki_group(
                    title="Ново барање за набавка!",
                    body=f"{naslov} (број {nalog}) – од {session['user']}",
                    url="/nabavki"
                )
                notify_novo_baranje(
                    nalog=nalog, naslov=naslov, kolicina=kolicina,
                    datum_itnost=datum_itnost, kreator=session["user"],
                )

            except ValueError:
                flash("Количина мора да биде број!", "error")
            except Exception as e:
                flash(f"Грешка при креирање: {str(e)}", "danger")
                conn.rollback()

        conn.close()
        return redirect(url_for("nabavki.nabavki", filter=request.args.get("filter", "all")))

    filter_type = request.args.get("filter", "all")
    if filter_type == "my_taken" and is_nabavki_user:
        rows = cursor.execute("""
            SELECT *, julianday(datum_itnost) - julianday(datum_kreiranje) AS days_diff
            FROM nabavki_requests WHERE prevzemeno_od=? ORDER BY nalog_broj DESC
        """, (session["user"],)).fetchall()
        page_title = "Мои превземени барања"
    elif is_nabavki_user:
        rows = cursor.execute("""
            SELECT *, julianday(datum_itnost) - julianday(datum_kreiranje) AS days_diff
            FROM nabavki_requests ORDER BY nalog_broj DESC
        """).fetchall()
        page_title = "Сите барања за набавки"
    else:
        rows = cursor.execute("""
            SELECT *, julianday(datum_itnost) - julianday(datum_kreiranje) AS days_diff
            FROM nabavki_requests WHERE username=? ORDER BY nalog_broj DESC
        """, (session["user"],)).fetchall()
        page_title = "Мои барања"

    requests_list = []
    for row in rows:
        r = dict(row)
        r["datum_kreiranje_formatted"] = _format_cet(r["datum_kreiranje"]) if r.get("datum_kreiranje") else "-"
        diff = r.get("days_diff")
        r["days_remaining"] = (
            f"{int(diff)} денови" if diff is not None and diff >= 0
            else ("Истечен рок!" if diff is not None else "Чека нарачување")
        )
        r["comments"] = cursor.execute("""
            SELECT user, comment, timestamp, slika FROM nabavki_comments
            WHERE req_id=? ORDER BY timestamp ASC
        """, (r["id"],)).fetchall()
        requests_list.append(r)

    users         = cursor.execute("SELECT username FROM users ORDER BY username").fetchall()
    nabavki_users = cursor.execute("""
        SELECT username FROM users
        WHERE user_group = 'Nabavki' OR is_admin = 1
        ORDER BY username
    """).fetchall()
    conn.close()

    return render_template("nabavki.html",
                           requests=requests_list,
                           is_admin=session.get("is_admin"),
                           is_nabavki_user=is_nabavki_user,
                           filter_type=filter_type,
                           page_title=page_title,
                           today=date.today().isoformat(),
                           users=users,
                           nabavki_users=nabavki_users)


# ─────────────────────────────────────────────────────────────
# DELETE SELECTED REQUESTS
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/delete_selected", methods=["POST"])
@login_required
def delete_selected():
    selected_ids = request.form.getlist("selected_ids")
    if not selected_ids:
        flash("Нема избрани барања за бришење!", "warning")
        return redirect(url_for("nabavki.nabavki"))
    try:
        conn         = get_db()
        cursor       = conn.cursor()
        placeholders = ",".join("?" for _ in selected_ids)
        cursor.execute(f"DELETE FROM nabavki_requests WHERE id IN ({placeholders})", selected_ids)
        count = cursor.rowcount
        conn.commit()
        flash(f"Успешно избришани {count} барања!", "success")
    except Exception as e:
        flash(f"Грешка при бришење: {str(e)}", "danger")
    finally:
        conn.close()
    return redirect(url_for("nabavki.nabavki"))


# ─────────────────────────────────────────────────────────────
# ADD COMMENT (со поддршка за слика)
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/add_comment/<int:req_id>", methods=["POST"])
@login_required
def add_comment(req_id):
    conn       = get_db()
    cursor     = conn.cursor()
    _ensure_comment_slika_column(cursor)

    comment    = request.form.get("comment", "").strip()
    slika_file = request.files.get("chat_slika")

    if not comment and (not slika_file or not slika_file.filename):
        flash("Коментарот е празен!", "error")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    req_row    = cursor.execute(
        "SELECT username, nalog_broj, naslov, prevzemeno_od FROM nabavki_requests WHERE id=?",
        (req_id,)
    ).fetchone()
    creator    = req_row["username"]      if req_row else None
    nalog      = req_row["nalog_broj"]    if req_row else "?"
    naslov     = req_row["naslov"]        if req_row else "?"
    prevzemeno = req_row["prevzemeno_od"] if req_row else None

    if not (session.get("is_admin") or session["user"] == creator or session.get("user_group") == "Nabavki"):
        flash("Немате дозвола да коментирате!", "error")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    try:
        chat_slika_filename = None
        if slika_file and slika_file.filename:
            save_dir      = os.path.join(current_app.config['STATIC_FOLDER'], "nabavki_chat")
            filename_base = f"chat_{req_id}_{int(time.time())}_{session['user']}"
            chat_slika_filename = _save_compressed_image(slika_file, save_dir, filename_base)
            if not chat_slika_filename:
                flash("Грешка при зачувување на сликата!", "warning")

        cursor.execute(
            "INSERT INTO nabavki_comments (req_id, user, comment, slika) VALUES (?,?,?,?)",
            (req_id, session["user"], comment, chat_slika_filename)
        )
        conn.commit()
        flash("Коментарот е успешно додаден!", "success")

        od = session["user"]
        if od == creator:
            if prevzemeno and prevzemeno != od:
                send_push_to_user(prevzemeno,
                    title=f"Нов коментар на барање {nalog}",
                    body=f"{od}: {comment[:80] if comment else '📷 Слика'}", url="/nabavki")
                notify_nov_komentar(nalog, naslov, comment or "[Слика]", od, prevzemeno)
        else:
            if creator and creator != od:
                send_push_to_user(creator,
                    title=f"Нов коментар на барање {nalog}",
                    body=f"{od}: {comment[:80] if comment else '📷 Слика'}", url="/nabavki")
                notify_nov_komentar(nalog, naslov, comment or "[Слика]", od, creator)

    except Exception as e:
        flash(f"Грешка: {str(e)}", "danger")
        conn.rollback()
    finally:
        conn.close()

    return redirect(url_for("nabavki.nabavki", filter=request.args.get("filter", "all")))


# ─────────────────────────────────────────────────────────────
# PREVZEMI REQUEST
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/prevzemi/<int:req_id>", methods=["POST"])
@login_required
def prevzemi(req_id):
    if not (session.get("user_group") == "Nabavki" or session.get("is_admin")):
        flash("Немате дозвола да превземате барања!", "danger")
        return redirect(url_for("nabavki.nabavki"))

    conn     = get_db()
    cursor   = conn.cursor()
    existing = cursor.execute(
        "SELECT prevzemeno_od, nalog_broj, username, naslov FROM nabavki_requests WHERE id=?",
        (req_id,)
    ).fetchone()

    if existing and existing["prevzemeno_od"]:
        flash(f'Ова барање веќе е превземено од {existing["prevzemeno_od"]}', "warning")
    else:
        try:
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            cursor.execute("""
                UPDATE nabavki_requests
                SET prevzemeno_od=?, datum_prevzemanje=?, status='Videno'
                WHERE id=?
            """, (session["user"], now, req_id))
            conn.commit()

            flash(f'Барањето е превземено од {session["user"]}!', "success")

            creator = existing["username"]
            nalog   = existing["nalog_broj"]
            naslov  = existing["naslov"]

            if creator and creator != session["user"]:
                send_push_to_user(creator,
                    title="Вашето барање е превземено",
                    body=f"Барање {nalog} го превзеде {session['user']}",
                    url="/nabavki")
            send_push_to_nabavki_group(
                title="Барање превземено",
                body=f"Барање {nalog} го превзеде {session['user']}",
                url="/nabavki")
            notify_promena_status(
                nalog=nalog, naslov=naslov,
                star_status="креирано", nov_status="Videno",
                kreator=creator, promeneto_od=session["user"],
            )
        except Exception as e:
            flash(f"Грешка: {str(e)}", "danger")

    conn.close()
    return redirect(url_for("nabavki.nabavki"))


# ─────────────────────────────────────────────────────────────
# TRANSFER REQUEST
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/transfer/<int:req_id>", methods=["POST"])
@login_required
def transfer(req_id):
    if not (session.get("user_group") == "Nabavki" or session.get("is_admin")):
        flash("Немате дозвола да префрлате барања!", "danger")
        return redirect(url_for("nabavki.nabavki"))

    new_user = request.form.get("novi_korisnik", "").strip()
    if not new_user:
        flash("Изберете корисник за префрлање!", "warning")
        return redirect(url_for("nabavki.nabavki"))

    conn   = get_db()
    cursor = conn.cursor()

    valid_user = cursor.execute("""
        SELECT username FROM users
        WHERE username = ? AND (user_group = 'Nabavki' OR is_admin = 1)
    """, (new_user,)).fetchone()
    if not valid_user:
        flash("Корисникот не постои или не е во групата Набавки!", "danger")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    existing = cursor.execute(
        "SELECT prevzemeno_od, nalog_broj, username, naslov FROM nabavki_requests WHERE id=?",
        (req_id,)
    ).fetchone()
    if not existing:
        flash("Барањето не постои!", "danger")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    try:
        now            = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        stari_korisnik = existing["prevzemeno_od"] or "—"
        nalog          = existing["nalog_broj"]
        creator        = existing["username"]
        naslov         = existing["naslov"]

        cursor.execute("""
            UPDATE nabavki_requests SET prevzemeno_od = ?, datum_prevzemanje = ? WHERE id = ?
        """, (new_user, now, req_id))
        conn.commit()

        komentar = f"Барањето е префрлено од {stari_korisnik} → {new_user} (од {session['user']})"
        cursor.execute(
            "INSERT INTO nabavki_comments (req_id, user, comment) VALUES (?, ?, ?)",
            (req_id, session["user"], komentar)
        )
        conn.commit()

        flash(f'Барањето {nalog} е префрлено на {new_user}!', "success")

        send_push_to_user(new_user,
            title="Ново префрлено барање!",
            body=f"Барање {nalog} ти е префрлено од {session['user']}",
            url="/nabavki")
        if stari_korisnik != "—" and stari_korisnik != new_user:
            send_push_to_user(stari_korisnik,
                title="Барањето ти е префрлено",
                body=f"Барање {nalog} е префрлено на {new_user} од {session['user']}",
                url="/nabavki")
        notify_nov_komentar(nalog, naslov, komentar, session["user"], new_user)

    except Exception as e:
        flash(f"Грешка при префрлање: {str(e)}", "danger")
        conn.rollback()
    finally:
        conn.close()

    return redirect(url_for("nabavki.nabavki", filter=request.args.get("filter", "all")))


# ─────────────────────────────────────────────────────────────
# UPDATE STATUS
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/update_status/<int:req_id>/<string:new_status>")
@login_required
def update_status(req_id, new_status):
    valid_statuses = {"Videno", "Naracano", "Dostaveno", "Zavrseno", "Prevzemeno"}
    if new_status not in valid_statuses:
        flash("Невалиден статус!", "danger")
        return redirect(url_for("nabavki.nabavki"))

    conn   = get_db()
    cursor = conn.cursor()
    try:
        request_data = cursor.execute("SELECT * FROM nabavki_requests WHERE id=?", (req_id,)).fetchone()
        if not request_data:
            flash("Барањето не постои!", "danger")
            conn.close()
            return redirect(url_for("nabavki.nabavki"))

        old_status    = request_data["status"]
        nalog         = request_data["nalog_broj"]
        naslov        = request_data["naslov"]
        creator       = request_data["username"]
        prevzemeno_od = request_data["prevzemeno_od"]

        cursor.execute("UPDATE nabavki_requests SET status=? WHERE id=?", (new_status, req_id))

        if new_status == "Zavrseno":
            max_row = cursor.execute("""
                SELECT MAX(CAST(SUBSTR(arhiva_broj, 4) AS INTEGER)) AS max_num
                FROM nabavki_archive WHERE arhiva_broj LIKE 'Arh%'
            """).fetchone()
            next_num    = (max_row["max_num"] or 0) + 1
            arhiva_broj = f"Arh{next_num:03d}"

            cursor.execute("""
                INSERT OR REPLACE INTO nabavki_archive
                (id, username, naslov, kolicina, opis, slika, datum_kreiranje, status,
                 datum_naracka, datum_priem, admin_naracal, datum_itnost, nalog_broj,
                 prevzemeno_od, datum_prevzemanje, arhivirano_od, arhiva_broj)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """, (request_data["id"], request_data["username"], request_data["naslov"],
                  request_data["kolicina"], request_data["opis"], request_data["slika"],
                  request_data["datum_kreiranje"], new_status, request_data["datum_naracka"],
                  request_data["datum_priem"], request_data["admin_naracal"], request_data["datum_itnost"],
                  request_data["nalog_broj"], request_data["prevzemeno_od"],
                  request_data["datum_prevzemanje"], session["user"], arhiva_broj))

            # ── Зачувај коментари пред бришење ──
            _ensure_comment_slika_column(cursor)
            _ensure_archive_comments_table(cursor)

            comments = cursor.execute("""
                SELECT user, comment, timestamp, slika FROM nabavki_comments
                WHERE req_id=? ORDER BY timestamp ASC
            """, (req_id,)).fetchall()

            for c in comments:
                slika_val = c["slika"] if "slika" in c.keys() else None
                cursor.execute("""
                    INSERT INTO nabavki_archive_comments
                        (archive_req_id, user, comment, slika, timestamp)
                    VALUES (?, ?, ?, ?, ?)
                """, (req_id, c["user"], c["comment"], slika_val, c["timestamp"]))

            cursor.execute("DELETE FROM nabavki_comments WHERE req_id=?", (req_id,))
            cursor.execute("DELETE FROM nabavki_requests WHERE id=?", (req_id,))

            flash(f'Барањето {nalog} е архивирано како {arhiva_broj}!', "success")

            send_push_to_nabavki_group(
                title="Барање архивирано",
                body=f"Барање {nalog} е завршено и архивирано.",
                url="/nabavki/arhiva")
        else:
            flash(f"Статусот е сменет на {new_status}!", "success")

        conn.commit()

        for recipient in {creator, prevzemeno_od} - {None, session["user"]}:
            send_push_to_user(recipient,
                title=f"Ажурирање на барање {nalog}",
                body=f"Статусот е сменет од {old_status} → {new_status}",
                url="/nabavki")

        notify_promena_status(
            nalog=nalog, naslov=naslov,
            star_status=old_status, nov_status=new_status,
            kreator=creator, promeneto_od=session["user"],
        )

    except Exception as e:
        flash(f"Грешка: {str(e)}", "danger")
        conn.rollback()
    finally:
        conn.close()

    return redirect(url_for("nabavki.nabavki"))


# ─────────────────────────────────────────────────────────────
# ARHIVA
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/arhiva")
@login_required
def arhiva():
    conn   = get_db()
    cursor = conn.cursor()
    _ensure_archive_comments_table(cursor)

    archived_raw = cursor.execute(
        "SELECT * FROM nabavki_archive ORDER BY arhivirano_na DESC"
    ).fetchall()

    archived = []
    for row in archived_raw:
        r = dict(row)
        r["comments"] = cursor.execute("""
            SELECT user, comment, slika, timestamp
            FROM nabavki_archive_comments
            WHERE archive_req_id = ?
            ORDER BY timestamp ASC
        """, (r["id"],)).fetchall()
        archived.append(r)

    conn.close()
    return render_template("nabavki_arhiva.html", archived_requests=archived)


# ─────────────────────────────────────────────────────────────
# EXPORT EXCEL
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/export/excel")
@login_required
def export_excel():
    conn   = get_db()
    cursor = conn.cursor()
    query  = "SELECT * FROM nabavki_requests WHERE 1=1"
    params = []
    if sf := request.args.get("status_filter"):
        query += " AND status=?"; params.append(sf)
    if do := request.args.get("datum_od"):
        query += " AND datum_kreiranje>=?"; params.append(do)
    if dt := request.args.get("datum_do"):
        query += " AND datum_kreiranje<=?"; params.append(dt + " 23:59:59")
    user_filter = request.args.get("user_filter") if session.get("is_admin") else session["user"]
    if user_filter:
        query += " AND username=?"; params.append(user_filter)
    query += " ORDER BY datum_kreiranje DESC"
    rows = cursor.execute(query, params).fetchall()
    conn.close()

    wb = Workbook(); ws = wb.active; ws.title = "Набавки"
    ws.append(["ID","Корисник","Наслов","Количина","Опис","Слика","Датум","Статус",
               "Датум нарачка","Датум прием","Admin нарачал"])
    for r in rows:
        ws.append([r["id"],r["username"],r["naslov"],r["kolicina"],
                   r["opis"] or "-",r["slika"] or "-",r["datum_kreiranje"],
                   r["status"],r["datum_naracka"] or "-",r["datum_priem"] or "-",
                   r["admin_naracal"] or "-"])
    output = BytesIO(); wb.save(output); output.seek(0)
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=f"nabavki_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")


# ─────────────────────────────────────────────────────────────
# EXPORT PDF
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/export/pdf")
@login_required
def export_pdf():
    conn   = get_db()
    cursor = conn.cursor()
    query  = "SELECT * FROM nabavki_requests WHERE 1=1"
    params = []
    if sf := request.args.get("status_filter"):
        query += " AND status=?"; params.append(sf)
    if do := request.args.get("datum_od"):
        query += " AND datum_kreiranje>=?"; params.append(do)
    if dt := request.args.get("datum_do"):
        query += " AND datum_kreiranje<=?"; params.append(dt + " 23:59:59")
    user_filter = request.args.get("user_filter") if session.get("is_admin") else session["user"]
    if user_filter:
        query += " AND username=?"; params.append(user_filter)
    query += " ORDER BY datum_kreiranje DESC"
    rows = cursor.execute(query, params).fetchall()
    conn.close()

    buffer = BytesIO()
    doc    = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    data   = [["ID","Корисник","Наслов","Кол.","Статус","Датум"]]
    for r in rows:
        data.append([r["id"],r["username"],r["naslov"],r["kolicina"],r["status"],r["datum_kreiranje"]])
    table = Table(data)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.grey),
        ("TEXTCOLOR",  (0,0), (-1,0), colors.whitesmoke),
        ("ALIGN",      (0,0), (-1,-1), "CENTER"),
        ("FONTNAME",   (0,0), (-1,0), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0,0), (-1,0), 12),
        ("BACKGROUND", (0,1), (-1,-1), colors.beige),
        ("GRID",       (0,0), (-1,-1), 1, colors.black),
    ]))
    doc.build([Paragraph("Листа на барања за набавки", styles["Heading1"]), table])
    buffer.seek(0)
    return send_file(buffer, mimetype="application/pdf", as_attachment=True,
                     download_name=f"nabavki_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf")


# ─────────────────────────────────────────────────────────────
# GET COMMENTS (враќа и слика)
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/comments/<int:req_id>")
@login_required
def get_comments(req_id):
    conn     = get_db()
    cursor   = conn.cursor()
    _ensure_comment_slika_column(cursor)
    comments = cursor.execute("""
        SELECT user, comment, timestamp, slika FROM nabavki_comments
        WHERE req_id=? ORDER BY timestamp ASC
    """, (req_id,)).fetchall()
    conn.close()
    return jsonify([dict(c) for c in comments])


# ─────────────────────────────────────────────────────────────
# API NABAVKI (за AJAX refresh)
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/api")
@login_required
def api():
    conn            = get_db()
    cursor          = conn.cursor()
    filter_type     = request.args.get("filter", "all")
    is_nabavki_user = session.get("is_admin") or session.get("user_group") == "Nabavki"

    if filter_type == "my_taken" and is_nabavki_user:
        rows = cursor.execute("""
            SELECT *, julianday(datum_itnost) - julianday(datum_kreiranje) AS days_diff
            FROM nabavki_requests WHERE prevzemeno_od=? ORDER BY nalog_broj DESC
        """, (session["user"],)).fetchall()
        page_title = "Мои превземени барања"
    elif is_nabavki_user:
        rows = cursor.execute("""
            SELECT *, julianday(datum_itnost) - julianday(datum_kreiranje) AS days_diff
            FROM nabavki_requests ORDER BY nalog_broj DESC
        """).fetchall()
        page_title = "Сите барања за набавки"
    else:
        rows = cursor.execute("""
            SELECT *, julianday(datum_itnost) - julianday(datum_kreiranje) AS days_diff
            FROM nabavki_requests WHERE username=? ORDER BY nalog_broj DESC
        """, (session["user"],)).fetchall()
        page_title = "Мои барања"

    requests_list = []
    for row in rows:
        r = dict(row)
        r["datum_kreiranje_formatted"] = _format_cet(r["datum_kreiranje"]) if r.get("datum_kreiranje") else "-"
        diff = r.get("days_diff")
        r["days_remaining"] = (
            f"{int(diff)} денови" if diff is not None and diff >= 0
            else ("Истечен рок!" if diff is not None else "Чека нарачување")
        )
        r["comments"] = cursor.execute("""
            SELECT user, comment, timestamp, slika FROM nabavki_comments
            WHERE req_id=? ORDER BY timestamp ASC
        """, (r["id"],)).fetchall()
        requests_list.append(r)

    conn.close()

    tbody_html = ""
    for r in requests_list:
        slika_cell = (
            f'<a href="/static/nabavki/{r["slika"]}" target="_blank">'
            f'<img src="/static/nabavki/{r["slika"]}" class="rounded shadow" style="max-width:90px;height:auto;"></a>'
            if r.get("slika") else '<span class="text-muted">Нема</span>'
        )
        status_map = {
            "Videno":    "bg-info",
            "Naracano":  "bg-warning text-dark",
            "Dostaveno": "bg-primary",
            "Zavrseno":  "bg-success",
            "Prevzemeno":"bg-success",
            "креирано":  "bg-secondary",
        }
        status_cls = status_map.get(r.get("status", "креирано"), "bg-secondary")
        prevzemeno = (
            f'<strong class="text-success">{r["prevzemeno_od"]}</strong>'
            + (f'<br><small class="text-muted">{r["datum_prevzemanje"]}</small>' if r.get("datum_prevzemanje") else "")
            if r.get("prevzemeno_od") else '<span class="text-muted">—</span>'
        )
        action_cell = '<span class="text-muted">—</span>'
        if session.get("is_admin") or session.get("user_group") == "Nabavki":
            options = "".join(
                f'<li><a class="dropdown-item status-change" href="#" data-status="{s}" data-id="{r["id"]}">{s}</a></li>'
                for s in ["Videno", "Naracano", "Dostaveno", "Zavrseno", "Prevzemeno"]
            )
            action_cell = f"""
                <div class="dropdown">
                    <button class="btn btn-outline-secondary btn-lg dropdown-toggle px-4 py-2"
                            type="button" data-bs-toggle="dropdown">Смени статус</button>
                    <ul class="dropdown-menu dropdown-menu-end shadow-lg">{options}</ul>
                </div>"""

        tbody_html += f"""
            <tr class="align-middle">
                <td><strong class="text-primary fs-5">{r.get("nalog_broj", "—")}</strong></td>
                <td>{r["datum_kreiranje_formatted"]}</td>
                <td>{r.get("username", "—")}</td>
                <td>{r.get("naslov", "—")}</td>
                <td class="fw-bold">{r.get("kolicina", 0)}</td>
                <td>{r.get("datum_itnost", "—")}</td>
                <td><span class="fw-bold text-success">{r["days_remaining"]}</span></td>
                <td>{slika_cell}</td>
                <td><span class="badge fs-6 py-2 px-4 rounded-pill {status_cls}">{r.get("status", "креирано").title()}</span></td>
                <td>{prevzemeno}</td>
                <td>
                    <button type="button" class="btn btn-outline-primary btn-lg px-4 py-2"
                            data-bs-toggle="modal" data-bs-target="#detailsModal{r.get('id')}">
                        <i class="fas fa-eye me-2"></i> Детали
                    </button>
                </td>
                <td>{action_cell}</td>
            </tr>"""

    return jsonify({
        "tbody_html": tbody_html,
        "page_title": page_title,
        "count":      len(requests_list),
    })
    # ─────────────────────────────────────────────────────────────
# ДОДАЈ СЛИКА НА ПОСТОЕЧКО БАРАЊЕ (ако нема слика)
# ─────────────────────────────────────────────────────────────
@nabavki_bp.route("/upload_slika/<int:req_id>", methods=["POST"])
@login_required
def upload_slika(req_id):
    conn   = get_db()
    cursor = conn.cursor()

    req_row = cursor.execute(
        "SELECT username, slika, nalog_broj FROM nabavki_requests WHERE id=?", (req_id,)
    ).fetchone()

    if not req_row:
        flash("Барањето не постои!", "danger")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    # Само креаторот, набавки групата или админот може да постави слика
    if not (session.get("is_admin") or
            session["user"] == req_row["username"] or
            session.get("user_group") == "Nabavki"):
        flash("Немате дозвола!", "danger")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    # Ако веќе има слика, не дозволи
    if req_row["slika"]:
        flash("Ова барање веќе има прикачена слика!", "warning")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    slika_file = request.files.get("nova_slika")
    if not slika_file or not slika_file.filename:
        flash("Нема избрана слика!", "warning")
        conn.close()
        return redirect(url_for("nabavki.nabavki"))

    try:
        save_dir      = os.path.join(current_app.config['STATIC_FOLDER'], "nabavki")
        filename_base = f"req_{req_id}_{int(time.time())}_{session['user']}"
        slika_filename = _save_compressed_image(slika_file, save_dir, filename_base)

        if not slika_filename:
            flash("Грешка при зачувување на сликата!", "danger")
            conn.close()
            return redirect(url_for("nabavki.nabavki"))

        cursor.execute(
            "UPDATE nabavki_requests SET slika=? WHERE id=?",
            (slika_filename, req_id)
        )
        conn.commit()
        flash(f"Сликата е успешно прикачена на барање {req_row['nalog_broj']}!", "success")

    except Exception as e:
        flash(f"Грешка: {str(e)}", "danger")
        conn.rollback()
    finally:
        conn.close()

    return redirect(url_for("nabavki.nabavki", filter=request.args.get("filter", "all")))