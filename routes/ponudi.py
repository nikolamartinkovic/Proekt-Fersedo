import json
import os
import re
import smtplib
import sqlite3
import time
from datetime import date, datetime
from email import encoders
from email.mime.base import MIMEBase
from email.mime.image import MIMEImage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr
from io import BytesIO

from flask import (
    Blueprint,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from utils.db import get_db
from utils.decorators import login_required, user_has_module

ponudi_bp = Blueprint("ponudi", __name__, url_prefix="/ponudi")


@ponudi_bp.before_request
def ensure_ponudi_access():
    if "user" not in session:
        return None
    endpoint = request.endpoint or ""
    archive_endpoints = {"ponudi.arhiva", "ponudi.export_excel", "ponudi.export_pdf"}
    required_module = "ponudi_arhiva" if endpoint in archive_endpoints else "ponudi"
    if user_has_module(required_module):
        return None
    flash("Немате дозвола за пристап до овој модул.", "danger")
    return redirect(url_for("auth.index"))


_IMG_MAX_SIZE = (1200, 1200)
_IMG_QUALITY = 72
STATUSI = ["Отворена", "Во преговори", "Прифатена", "Одбиена", "Завршена"]
DEFAULT_DRAFT_TITLE = "Недовршена понуда"


def _is_manager():
    return bool(session.get("is_admin") or session.get("user_group") == "Nabavki")


def _can_comment(ponuda_row):
    return bool(
        session.get("is_admin")
        or session.get("user_group") == "Nabavki"
        or (ponuda_row and ponuda_row["username"] == session.get("user"))
    )


def _save_compressed_image(file_storage, save_dir, filename_base):
    try:
        os.makedirs(save_dir, exist_ok=True)
        final_name = f"{filename_base}.jpg"
        save_path = os.path.join(save_dir, final_name)
        img = PILImage.open(file_storage)
        if img.mode in ("RGBA", "P", "LA"):
            img = img.convert("RGB")
        img.thumbnail(_IMG_MAX_SIZE, PILImage.Resampling.LANCZOS)
        img.save(save_path, format="JPEG", quality=_IMG_QUALITY, optimize=True)
        return final_name
    except Exception as exc:
        print(f"[PONUDI IMAGE] {exc}")
        return None


def _column_names(cursor, table_name):
    cols = cursor.execute(f"PRAGMA table_info({table_name})").fetchall()
    names = set()
    for col in cols:
        try:
            names.add(col["name"])
        except Exception:
            names.add(col[1])
    return names


def _ensure_column(cursor, table_name, column_name, definition):
    if column_name not in _column_names(cursor, table_name):
        cursor.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}")


def _seed_existing_suppliers(cursor):
    existing_names = {
        (row["naziv"] or "").strip().lower()
        for row in cursor.execute("SELECT naziv FROM ponudi_dobavuvaci").fetchall()
    }
    for table_name in ("ponudi", "ponudi_archive"):
        rows = cursor.execute(
            f"SELECT DISTINCT dobavuvac FROM {table_name} WHERE dobavuvac IS NOT NULL AND dobavuvac != ''"
        ).fetchall()
        for row in rows:
            raw_value = (row["dobavuvac"] or "").strip()
            if not raw_value:
                continue
            for part in re.split(r"[;,]", raw_value):
                naziv = part.strip()
                if not naziv:
                    continue
                key = naziv.lower()
                if key in existing_names:
                    continue
                cursor.execute(
                    """
                    INSERT OR IGNORE INTO ponudi_dobavuvaci
                    (naziv, kontakt_lice, email, telefon, adresa, zabeleska, aktiven)
                    VALUES (?, '', '', '', '', '', 1)
                    """,
                    (naziv,),
                )
                existing_names.add(key)


def _ensure_tables(cursor):
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi (
            id               INTEGER PRIMARY KEY AUTOINCREMENT,
            ponuda_broj      TEXT UNIQUE,
            username         TEXT NOT NULL,
            naslov           TEXT NOT NULL,
            dobavuvac        TEXT NOT NULL DEFAULT '',
            cena             REAL,
            valuta           TEXT DEFAULT 'EUR',
            rok_isporaka     TEXT,
            opis             TEXT,
            slika            TEXT,
            status           TEXT DEFAULT 'Отворена',
            datum_kreiranje  TEXT DEFAULT CURRENT_TIMESTAMP,
            datum_vaznost    TEXT,
            arhivirano_od    TEXT,
            arhivirano_na    TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_comments (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            ponuda_id  INTEGER NOT NULL,
            user       TEXT,
            comment    TEXT,
            slika      TEXT,
            timestamp  TEXT DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_archive (
            id               INTEGER PRIMARY KEY,
            ponuda_broj      TEXT,
            username         TEXT,
            naslov           TEXT,
            dobavuvac        TEXT,
            cena             REAL,
            valuta           TEXT,
            rok_isporaka     TEXT,
            opis             TEXT,
            slika            TEXT,
            status           TEXT,
            datum_kreiranje  TEXT,
            datum_vaznost    TEXT,
            arhivirano_od    TEXT,
            arhivirano_na    TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_archive_comments (
            id                 INTEGER PRIMARY KEY AUTOINCREMENT,
            archive_ponuda_id  INTEGER NOT NULL,
            user               TEXT,
            comment            TEXT,
            slika              TEXT,
            timestamp          TEXT
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS pon_sequence (
            id INTEGER PRIMARY KEY,
            last_num INTEGER DEFAULT 0
        )
        """
    )
    cursor.execute("INSERT OR IGNORE INTO pon_sequence (id, last_num) VALUES (1, 0)")
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_dobavuvaci (
            id            INTEGER PRIMARY KEY AUTOINCREMENT,
            naziv         TEXT NOT NULL UNIQUE,
            kontakt_lice  TEXT,
            email         TEXT,
            telefon       TEXT,
            adresa        TEXT,
            zabeleska     TEXT,
            aktiven       INTEGER DEFAULT 1,
            created_at    TEXT DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_supplier_links (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            ponuda_id       INTEGER NOT NULL,
            supplier_id     INTEGER NOT NULL,
            supplier_name   TEXT NOT NULL,
            supplier_email  TEXT,
            created_at      TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(ponuda_id, supplier_id)
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_items (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            ponuda_id    INTEGER NOT NULL,
            item_name    TEXT NOT NULL,
            quantity     REAL DEFAULT 1,
            comment      TEXT,
            image        TEXT,
            sort_order   INTEGER DEFAULT 0,
            created_at   TEXT DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS ponudi_archive_items (
            id                 INTEGER PRIMARY KEY AUTOINCREMENT,
            archive_ponuda_id  INTEGER NOT NULL,
            item_name          TEXT NOT NULL,
            quantity           REAL DEFAULT 1,
            comment            TEXT,
            image              TEXT,
            sort_order         INTEGER DEFAULT 0,
            created_at         TEXT
        )
        """
    )
    _ensure_column(cursor, "ponudi", "kolicina", "REAL DEFAULT 1")
    _ensure_column(cursor, "ponudi", "is_draft", "INTEGER DEFAULT 0")
    _ensure_column(cursor, "ponudi", "updated_at", "TEXT")
    _ensure_column(cursor, "ponudi_archive", "kolicina", "REAL DEFAULT 1")
    _ensure_column(cursor, "ponudi_archive", "updated_at", "TEXT")
    _seed_existing_suppliers(cursor)


def _next_ponuda_broj(cursor):
    cursor.execute("UPDATE pon_sequence SET last_num = last_num + 1 WHERE id = 1")
    row = cursor.execute("SELECT last_num FROM pon_sequence WHERE id = 1").fetchone()
    return f"Pon{row['last_num']:03d}"


def _fmt_date(value):
    if not value:
        return "—"
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").strftime("%d-%m-%Y")
    except Exception:
        return str(value)


def _fmt_datetime(value):
    if not value:
        return "—"
    text = str(value).strip()
    for raw_value, fmt in ((text[:19], "%Y-%m-%d %H:%M:%S"), (text[:16], "%Y-%m-%d %H:%M")):
        try:
            return datetime.strptime(raw_value, fmt).strftime("%d-%m-%Y %H:%M")
        except Exception:
            continue
    return text


def _safe_filename_part(value):
    text = re.sub(r"[^A-Za-z0-9А-Ша-шЀ-џ_-]+", "_", str(value or "").strip())
    return text.strip("_") or "ponuda"


def _company_profile():
    cfg = current_app.config
    city_line = ", ".join(part for part in [cfg.get("COMPANY_ADDRESS", ""), cfg.get("COMPANY_CITY", "")] if part)
    return {
        "name": cfg.get("COMPANY_NAME", "Fersedo"),
        "address_line": city_line,
        "phone": cfg.get("COMPANY_PHONE", ""),
        "email": cfg.get("COMPANY_EMAIL", ""),
        "logo_path": os.path.join(cfg["STATIC_FOLDER"], "logo2.png"),
    }


def _mail_sender_profile():
    cfg = current_app.config
    email_user = cfg.get("EMAIL_HOST_USER", "").strip()
    return {
        "host": cfg.get("EMAIL_HOST"),
        "port": cfg.get("EMAIL_PORT"),
        "user": email_user,
        "password": cfg.get("EMAIL_HOST_PASSWORD"),
        "from_header": formataddr((cfg.get("EMAIL_FROM_NAME", "Info Fersedo"), email_user)) if email_user else cfg.get("EMAIL_FROM_NAME", "Info Fersedo"),
    }


def _offer_email_subject(ponuda):
    return f"Барање за понуда {ponuda.get('ponuda_broj') or ''} - Fersedo".strip(" -")


def _build_offer_email_html(ponuda, supplier):
    company = _company_profile()
    creator_name = ponuda.get("username") or session.get("user", "Fersedo")
    creator_email = (ponuda.get("creator_email") or "").strip()
    supplier_name = supplier.get("naziv") or supplier.get("supplier_name") or "добавувач"
    items = ponuda.get("items") or _fallback_items_from_ponuda(ponuda)

    item_rows = []
    for index, item in enumerate(items, start=1):
        item_rows.append(
            f"""
            <tr>
                <td style="padding:10px;border:1px solid #dbe3f0;text-align:center;">{index}</td>
                <td style="padding:10px;border:1px solid #dbe3f0;">{item.get("item_name") or "—"}</td>
                <td style="padding:10px;border:1px solid #dbe3f0;text-align:center;">{item.get("quantity") or "—"}</td>
                <td style="padding:10px;border:1px solid #dbe3f0;">{item.get("comment") or "—"}</td>
            </tr>
            """
        )

    contact_lines = [company.get("name", "Fersedo")]
    if creator_email:
        contact_lines.append(f"Контакт email: {creator_email}")
    if company.get("phone"):
        contact_lines.append(f"Телефон: {company['phone']}")
    contact_block = "<br>".join(contact_lines)

    return f"""
    <html>
    <body style="margin:0;padding:0;background-color:#f4f7fb;">
        <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%" style="background-color:#f4f7fb;margin:0;padding:0;">
            <tr>
                <td align="center" style="padding:24px 12px;">
                    <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="760" style="width:760px;max-width:760px;background-color:#ffffff;border:1px solid #dbe3f0;">
                        <tr>
                            <td style="background-color:#1e3168;padding:20px 28px;">
                                <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%">
                                    <tr>
                                        <td valign="middle" style="font-family:Arial,Helvetica,sans-serif;color:#ffffff;">
                                            <div style="font-size:16px;line-height:20px;font-weight:bold;margin:0 0 8px 0;">Барање за понуда</div>
                                            <div style="font-size:14px;line-height:20px;color:#dbe7ff;">Број: {ponuda.get('ponuda_broj') or '—'} | Датум: {_fmt_date(date.today().isoformat())}</div>
                                        </td>
                                        <td valign="middle" align="right" width="180">
                                            <img src="cid:fersedo-logo" alt="Fersedo" width="150" style="display:block;border:0;outline:none;text-decoration:none;width:150px;height:auto;">
                                        </td>
                                    </tr>
                                </table>
                            </td>
                        </tr>
                        <tr>
                            <td style="padding:28px;font-family:Arial,Helvetica,sans-serif;color:#0f172a;">
                                <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%">
                                    <tr>
                                        <td style="font-size:16px;line-height:24px;padding-bottom:14px;">Почитувани {supplier_name},</td>
                                    </tr>
                                    <tr>
                                        <td style="font-size:15px;line-height:24px;padding-bottom:10px;">Во прилог Ви испраќаме барање за понуда од <strong>{company.get('name', 'Fersedo')}</strong>.</td>
                                    </tr>
                                    <tr>
                                        <td style="font-size:15px;line-height:24px;padding-bottom:18px;">Ве молиме да ни доставите Ваш одговор за наведените артикли со соодветна цена, рок и услови според Вашата понуда.</td>
                                    </tr>
                                </table>

                                <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%" style="border-collapse:collapse;margin:0 0 22px 0;">
                                    <tr style="background-color:#0f172a;color:#ffffff;">
                                        <th align="center" style="padding:10px 8px;border:1px solid #dbe3f0;font-family:Arial,Helvetica,sans-serif;font-size:13px;line-height:18px;">Р.б.</th>
                                        <th align="left" style="padding:10px 8px;border:1px solid #dbe3f0;font-family:Arial,Helvetica,sans-serif;font-size:13px;line-height:18px;">Артикл / предмет</th>
                                        <th align="center" style="padding:10px 8px;border:1px solid #dbe3f0;font-family:Arial,Helvetica,sans-serif;font-size:13px;line-height:18px;">Количина</th>
                                        <th align="left" style="padding:10px 8px;border:1px solid #dbe3f0;font-family:Arial,Helvetica,sans-serif;font-size:13px;line-height:18px;">Коментар</th>
                                    </tr>
                                    {''.join(item_rows)}
                                </table>

                                <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%" style="background-color:#f8fafc;border:1px solid #dbe3f0;">
                                    <tr>
                                        <td style="padding:16px 18px;font-family:Arial,Helvetica,sans-serif;">
                                            <div style="font-size:14px;line-height:18px;font-weight:bold;color:#0f172a;margin-bottom:6px;">Контакт за одговор</div>
                                            <div style="font-size:14px;line-height:22px;color:#334155;">{contact_block}</div>
                                        </td>
                                    </tr>
                                </table>

                                <table role="presentation" cellpadding="0" cellspacing="0" border="0" width="100%" style="margin-top:22px;">
                                    <tr>
                                        <td style="font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;color:#0f172a;padding-bottom:6px;">Ви благодариме однапред.</td>
                                    </tr>
                                    <tr>
                                        <td style="font-family:Arial,Helvetica,sans-serif;font-size:15px;line-height:24px;color:#0f172a;">Со почит,<br><strong>{creator_name}</strong><br>{company.get('name', 'Fersedo')}</td>
                                    </tr>
                                </table>
                            </td>
                        </tr>
                    </table>
                </td>
            </tr>
        </table>
    </body>
    </html>
    """


def _send_offer_email(ponuda, supplier):
    supplier_email = (supplier.get("email") or supplier.get("supplier_email") or "").strip()
    if not supplier_email:
        return {"success": False, "message": "Добавувачот нема внесен email."}

    sender = _mail_sender_profile()
    if not sender["host"] or not sender["port"] or not sender["user"] or not sender["password"]:
        return {"success": False, "message": "Email конфигурацијата не е наместена."}

    supplier_name = supplier.get("naziv") or supplier.get("supplier_name") or "dobavuvac"
    subject = _offer_email_subject(ponuda)
    html_body = _build_offer_email_html(ponuda, supplier)
    pdf_buffer = _generate_offer_pdf_v2(ponuda, supplier)
    filename = f"{_safe_filename_part(ponuda.get('ponuda_broj') or ponuda.get('naslov'))}_{_safe_filename_part(supplier_name)}.pdf"

    msg = MIMEMultipart()
    msg["From"] = sender["from_header"]
    msg["To"] = supplier_email
    msg["Subject"] = subject
    msg.attach(MIMEText(html_body, "html", "utf-8"))

    logo_path = _company_profile().get("logo_path", "")
    if logo_path and os.path.exists(logo_path):
        with open(logo_path, "rb") as logo_file:
            inline_logo = MIMEImage(logo_file.read())
        inline_logo.add_header("Content-ID", "<fersedo-logo>")
        inline_logo.add_header("Content-Disposition", "inline", filename=os.path.basename(logo_path))
        msg.attach(inline_logo)

    attachment = MIMEBase("application", "pdf")
    attachment.set_payload(pdf_buffer.read())
    encoders.encode_base64(attachment)
    attachment.add_header("Content-Disposition", f'attachment; filename="{filename}"')
    msg.attach(attachment)

    try:
        with smtplib.SMTP(sender["host"], sender["port"], timeout=20) as server:
            server.ehlo()
            server.starttls()
            server.ehlo()
            server.login(sender["user"], sender["password"])
            server.sendmail(sender["user"], [supplier_email], msg.as_string())
        return {"success": True, "message": f"Понудата е испратена до {supplier_email}."}
    except Exception as exc:
        current_app.logger.warning("Offer email send failed for %s: %s", supplier_email, exc)
        return {"success": False, "message": str(exc)}


def _generate_offer_pdf(ponuda, supplier):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=32, rightMargin=32, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    regular_font = "DejaVuSans"
    bold_font = "DejaVuSans-Bold"

    title_style = ParagraphStyle("OfferTitle", parent=styles["Heading1"], fontName=bold_font, fontSize=18, leading=22, textColor=colors.HexColor("#0f172a"), spaceAfter=8)
    label_style = ParagraphStyle("OfferLabel", parent=styles["Normal"], fontName=bold_font, fontSize=9, leading=12, textColor=colors.HexColor("#334155"), spaceAfter=2)
    body_style = ParagraphStyle("OfferBody", parent=styles["Normal"], fontName=regular_font, fontSize=10, leading=13, textColor=colors.HexColor("#111827"))

    company = _company_profile()
    creator_email = (ponuda.get("creator_email") or "").strip()
    elements = []
    logo_path = company["logo_path"]
    logo = RLImage(logo_path, width=110, height=36) if os.path.exists(logo_path) else None

    left_header = []
    if logo:
        left_header.append(logo)
        left_header.append(Spacer(1, 8))
    if not logo:
        left_header.append(Paragraph(company["name"], ParagraphStyle("CompanyName", parent=title_style, fontSize=16, leading=20)))
    left_header.extend([
        Paragraph("Барање за понуда", title_style),
        Paragraph(f"Број: {ponuda.get('ponuda_broj') or '—'}", body_style),
        Paragraph(f"Датум: {_fmt_date(date.today().isoformat())}", body_style),
    ])
    if company["address_line"]:
        left_header.append(Paragraph(company["address_line"], body_style))
    if company["phone"]:
        left_header.append(Paragraph(f"Телефон: {company['phone']}", body_style))
    if creator_email:
        left_header.append(Paragraph(f"Контакт email: {creator_email}", body_style))

    supplier_lines = [
        Paragraph("До добавувач", label_style),
        Paragraph(supplier.get("naziv") or supplier.get("supplier_name") or "—", ParagraphStyle("SupplierName", parent=body_style, fontName=bold_font, fontSize=12, leading=15)),
    ]
    if supplier.get("kontakt_lice"):
        supplier_lines.append(Paragraph(f"Контакт лице: {supplier['kontakt_lice']}", body_style))
    if supplier.get("adresa"):
        supplier_lines.append(Paragraph(supplier["adresa"], body_style))
    if supplier.get("telefon"):
        supplier_lines.append(Paragraph(f"Телефон: {supplier['telefon']}", body_style))
    if supplier.get("email"):
        supplier_lines.append(Paragraph(f"Email: {supplier['email']}", body_style))

    header_table = Table([[left_header, supplier_lines]], colWidths=[doc.width * 0.56, doc.width * 0.44], hAlign="LEFT")
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#f8fafc")), ("BOX", (1, 0), (1, 0), 0.8, colors.HexColor("#cbd5e1")), ("PADDING", (0, 0), (-1, -1), 10)]))
    elements.append(header_table)
    elements.append(Spacer(1, 18))
    elements.append(Paragraph("Ве молиме доставете понуда за наведениот артикл/предмет подолу. Документот е подготвен од Fersedo како барање за понуда.", body_style))
    elements.append(Spacer(1, 14))

    article_table = Table([["Р.б.", "Артикл / предмет на барање", "Количина"], ["1", ponuda.get("naslov") or "—", str(ponuda.get("kolicina") or "—")]], colWidths=[42, doc.width - 42 - 96, 96], repeatRows=1, hAlign="LEFT")
    article_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), bold_font), ("FONTNAME", (0, 1), (-1, -1), regular_font), ("FONTSIZE", (0, 0), (-1, -1), 10), ("LEADING", (0, 0), (-1, -1), 13), ("GRID", (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")), ("BACKGROUND", (0, 1), (-1, -1), colors.white), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, -1), "CENTER"), ("ALIGN", (2, 0), (2, -1), "CENTER"), ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    elements.append(article_table)
    elements.append(Spacer(1, 18))

    footer_table = Table([[Paragraph("Подготвил", label_style), Paragraph("Добавувач", label_style)], [Paragraph(session.get("user", "Fersedo"), body_style), Paragraph(supplier.get("naziv") or supplier.get("supplier_name") or "—", body_style)]], colWidths=[doc.width / 2, doc.width / 2], hAlign="LEFT")
    footer_table.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.HexColor("#cbd5e1")), ("TOPPADDING", (0, 0), (-1, -1), 8)]))
    elements.append(footer_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def _supplier_rows(cursor, only_active=False):
    query = "SELECT * FROM ponudi_dobavuvaci"
    if only_active:
        query += " WHERE aktiven = 1"
    query += " ORDER BY naziv COLLATE NOCASE"
    return [dict(row) for row in cursor.execute(query).fetchall()]


def _supplier_map(cursor):
    return {row["id"]: dict(row) for row in cursor.execute("SELECT * FROM ponudi_dobavuvaci").fetchall()}


def _user_email(cursor, username):
    if not username:
        return ""
    row = cursor.execute("SELECT email FROM users WHERE username = ?", (username,)).fetchone()
    return (row["email"] or "").strip() if row else ""


def _parse_supplier_ids(values):
    result = []
    for value in values:
        try:
            supplier_id = int(value)
        except (TypeError, ValueError):
            continue
        if supplier_id not in result:
            result.append(supplier_id)
    return result


def _get_links_for_ponuda(cursor, ponuda_id):
    rows = cursor.execute(
        """
        SELECT supplier_id, supplier_name, supplier_email
        FROM ponudi_supplier_links
        WHERE ponuda_id = ?
        ORDER BY id ASC
        """,
        (ponuda_id,),
    ).fetchall()
    return [dict(row) for row in rows]


def _fallback_items_from_ponuda(record):
    return [
        {
            "item_name": record.get("naslov") or "—",
            "quantity": record.get("kolicina") or 1,
            "comment": "",
            "image": record.get("slika") or "",
            "sort_order": 0,
        }
    ]


def _item_image_relative_path(filename):
    if not filename:
        return ""
    for folder in ("ponudi_items", "ponudi"):
        fs_path = os.path.join(current_app.config["STATIC_FOLDER"], folder, filename)
        if os.path.exists(fs_path):
            return f"{folder}/{filename}"
    return ""


def _item_image_filesystem_path(filename):
    if not filename:
        return ""
    for folder in ("ponudi_items", "ponudi"):
        fs_path = os.path.join(current_app.config["STATIC_FOLDER"], folder, filename)
        if os.path.exists(fs_path):
            return fs_path
    return ""


def _get_items_for_ponuda(cursor, ponuda_id, archived=False, record=None):
    table_name = "ponudi_archive_items" if archived else "ponudi_items"
    id_col = "archive_ponuda_id" if archived else "ponuda_id"
    rows = cursor.execute(
        f"""
        SELECT item_name, quantity, comment, image, sort_order
        FROM {table_name}
        WHERE {id_col} = ?
        ORDER BY sort_order ASC, id ASC
        """,
        (ponuda_id,),
    ).fetchall()
    items = [dict(row) for row in rows]
    if items:
        return items
    if record:
        return _fallback_items_from_ponuda(record)
    return []


def _replace_items_for_ponuda(cursor, ponuda_id, items):
    cursor.execute("DELETE FROM ponudi_items WHERE ponuda_id = ?", (ponuda_id,))
    for index, item in enumerate(items):
        cursor.execute(
            """
            INSERT INTO ponudi_items (ponuda_id, item_name, quantity, comment, image, sort_order)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                ponuda_id,
                item["item_name"],
                item["quantity"],
                item.get("comment", ""),
                item.get("image", ""),
                index,
            ),
        )


def _copy_items_to_archive(cursor, ponuda_id):
    items = cursor.execute(
        """
        SELECT item_name, quantity, comment, image, sort_order, created_at
        FROM ponudi_items
        WHERE ponuda_id = ?
        ORDER BY sort_order ASC, id ASC
        """,
        (ponuda_id,),
    ).fetchall()
    for item in items:
        cursor.execute(
            """
            INSERT INTO ponudi_archive_items
            (archive_ponuda_id, item_name, quantity, comment, image, sort_order, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ponuda_id,
                item["item_name"],
                item["quantity"],
                item["comment"],
                item["image"],
                item["sort_order"],
                item["created_at"],
            ),
        )


def _parse_item_rows(req):
    names = req.form.getlist("item_name[]")
    quantities = req.form.getlist("item_quantity[]")
    comments = req.form.getlist("item_comment[]")
    existing_images = req.form.getlist("existing_item_image[]")
    image_files = req.files.getlist("item_image[]")

    total_rows = max(len(names), len(quantities), len(comments), len(existing_images), len(image_files))
    items = []
    save_dir = os.path.join(current_app.config["STATIC_FOLDER"], "ponudi_items")

    for index in range(total_rows):
        item_name = names[index].strip() if index < len(names) and names[index] else ""
        quantity_raw = quantities[index].strip() if index < len(quantities) and quantities[index] else ""
        comment = comments[index].strip() if index < len(comments) and comments[index] else ""
        existing_image = existing_images[index].strip() if index < len(existing_images) and existing_images[index] else ""
        image_file = image_files[index] if index < len(image_files) else None

        if not item_name and not quantity_raw and not comment and not existing_image and not (image_file and image_file.filename):
            continue

        quantity = None
        if quantity_raw:
            try:
                quantity = float(quantity_raw.replace(",", "."))
            except ValueError:
                raise ValueError(f"Количината за артикл {index + 1} мора да биде број.")
        if quantity is None or quantity <= 0:
            raise ValueError(f"Количината за артикл {index + 1} мора да биде поголема од 0.")
        if not item_name:
            raise ValueError(f"Називот за артикл {index + 1} е задолжителен.")

        image_name = existing_image
        if image_file and image_file.filename:
            filename_base = f"pon_item_{int(time.time())}_{session['user']}_{index}"
            saved_name = _save_compressed_image(image_file, save_dir, filename_base)
            if saved_name:
                image_name = saved_name

        items.append(
            {
                "item_name": item_name,
                "quantity": quantity,
                "comment": comment,
                "image": image_name,
            }
        )

    return items


def _supplier_names_from_text(raw_value):
    if not raw_value:
        return []
    return [part.strip() for part in re.split(r"[;,]", raw_value) if part.strip()]


def _sync_supplier_links(cursor, ponuda_id, supplier_ids, suppliers):
    cursor.execute("DELETE FROM ponudi_supplier_links WHERE ponuda_id = ?", (ponuda_id,))
    selected_names = []
    for supplier_id in supplier_ids:
        supplier = suppliers.get(supplier_id)
        if not supplier:
            continue
        cursor.execute(
            """
            INSERT OR IGNORE INTO ponudi_supplier_links
            (ponuda_id, supplier_id, supplier_name, supplier_email)
            VALUES (?, ?, ?, ?)
            """,
            (ponuda_id, supplier_id, supplier["naziv"], supplier.get("email")),
        )
        selected_names.append(supplier["naziv"])
    joined_names = ", ".join(selected_names)
    cursor.execute("UPDATE ponudi SET dobavuvac = ? WHERE id = ?", (joined_names, ponuda_id))
    return joined_names


def _decorate_ponuda(cursor, row, include_comments=True):
    record = dict(row)
    links = _get_links_for_ponuda(cursor, record["id"])
    record["supplier_links"] = links
    record["supplier_ids"] = [item["supplier_id"] for item in links]
    record["items"] = _get_items_for_ponuda(cursor, record["id"], archived=False, record=record)
    for item in record["items"]:
        item["image_url"] = _item_image_relative_path(item.get("image"))
    record["kolicina_total"] = sum((item.get("quantity") or 0) for item in record["items"])
    record["primary_image"] = next(
        (item.get("image_url") for item in record["items"] if item.get("image_url")),
        _item_image_relative_path(record.get("slika") or ""),
    )
    record["supplier_display"] = ", ".join(item["supplier_name"] for item in links) if links else ", ".join(
        _supplier_names_from_text(record.get("dobavuvac", ""))
    )
    record["datum_kreiranje_display"] = _fmt_datetime(record.get("datum_kreiranje"))
    record["updated_at_display"] = _fmt_datetime(record.get("updated_at") or record.get("datum_kreiranje"))
    if include_comments:
        record["comments"] = cursor.execute(
            """
            SELECT user, comment, slika, timestamp
            FROM ponudi_comments
            WHERE ponuda_id = ?
            ORDER BY timestamp ASC
            """,
            (record["id"],),
        ).fetchall()
    return record


def _load_editable_draft(cursor, draft_id):
    if not draft_id:
        return None
    row = cursor.execute("SELECT * FROM ponudi WHERE id = ? AND is_draft = 1", (draft_id,)).fetchone()
    if not row:
        return None
    if _is_manager() or row["username"] == session.get("user"):
        return row
    return None


def _save_or_update_ponuda(cursor, action, draft_row, slika_filename, title, kolicina, chat_comment, items):
    is_draft = 1 if action == "save_draft" else 0
    total_quantity = sum((item.get("quantity") or 0) for item in items) if items else (kolicina or 0)
    status = "Draft" if is_draft else "Отворена"

    if draft_row:
        final_image = slika_filename or draft_row["slika"]
        cursor.execute(
            """
            UPDATE ponudi
            SET naslov = ?, kolicina = ?, slika = ?, status = ?, is_draft = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (title, total_quantity, final_image, status, is_draft, draft_row["id"]),
        )
        ponuda_id = draft_row["id"]
        ponuda_broj = draft_row["ponuda_broj"]
    else:
        cursor.execute(
            """
            INSERT INTO ponudi
            (username, naslov, dobavuvac, cena, valuta, rok_isporaka, opis, slika, status,
             datum_kreiranje, datum_vaznost, kolicina, is_draft, updated_at)
            VALUES (?, ?, '', NULL, NULL, NULL, NULL, ?, ?, CURRENT_TIMESTAMP, NULL, ?, ?, CURRENT_TIMESTAMP)
            """,
            (session["user"], title, slika_filename, status, total_quantity, is_draft),
        )
        ponuda_id = cursor.lastrowid
        ponuda_broj = None

    if not is_draft and not ponuda_broj:
        ponuda_broj = _next_ponuda_broj(cursor)
        cursor.execute("UPDATE ponudi SET ponuda_broj = ? WHERE id = ?", (ponuda_broj, ponuda_id))

    if chat_comment:
        cursor.execute(
            """
            INSERT INTO ponudi_comments (ponuda_id, user, comment)
            VALUES (?, ?, ?)
            """,
            (ponuda_id, session["user"], chat_comment),
        )

    _replace_items_for_ponuda(cursor, ponuda_id, items)
    return ponuda_id, ponuda_broj


def _generate_offer_pdf_v2(ponuda, supplier):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=34, rightMargin=34, topMargin=18, bottomMargin=32)
    styles = getSampleStyleSheet()
    regular_font = "DejaVuSans"
    bold_font = "DejaVuSans-Bold"

    title_style = ParagraphStyle(
        "OfferTitleV2",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=8,
    )
    title_left_style = ParagraphStyle(
        "OfferTitleLeftV2",
        parent=title_style,
        alignment=0,
        leftIndent=8,
        spaceAfter=10,
    )
    label_style = ParagraphStyle(
        "OfferLabelV2",
        parent=styles["Normal"],
        fontName=bold_font,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
        spaceAfter=2,
    )
    body_style = ParagraphStyle(
        "OfferBodyV2",
        parent=styles["Normal"],
        fontName=regular_font,
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#111827"),
    )
    small_style = ParagraphStyle(
        "OfferSmallV2",
        parent=body_style,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
    )
    table_cell_style = ParagraphStyle(
        "OfferTableCellV2",
        parent=body_style,
        fontSize=9,
        leading=11,
        wordWrap="LTR",
    )
    table_center_style = ParagraphStyle(
        "OfferTableCenterV2",
        parent=table_cell_style,
        alignment=1,
    )

    company = _company_profile()
    creator_email = (ponuda.get("creator_email") or "").strip()
    items = ponuda.get("items") or _fallback_items_from_ponuda(ponuda)
    elements = []

    logo_path = company["logo_path"]
    logo = RLImage(logo_path, width=110, height=36) if os.path.exists(logo_path) else None

    if logo:
        logo_table = Table([[logo]], colWidths=[doc.width], hAlign="LEFT")
        logo_table.setStyle(
            TableStyle(
                [
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ]
            )
        )
        elements.append(logo_table)
        elements.append(Spacer(1, 10))

    left_header = [Spacer(1, 30)]
    left_header.extend(
        [
            Paragraph("Барање за понуда", title_left_style),
            Spacer(1, 4),
            Paragraph(f"Број: {ponuda.get('ponuda_broj') or '—'}", body_style),
            Paragraph(f"Датум: {_fmt_date(date.today().isoformat())}", body_style),
        ]
    )
    if creator_email:
        left_header.append(Paragraph(f"Контакт email: {creator_email}", body_style))

    supplier_lines = [
        Spacer(1, 28),
        Paragraph("До добавувач", label_style),
        Paragraph(
            supplier.get("naziv") or supplier.get("supplier_name") or "—",
            ParagraphStyle("SupplierNameV2", parent=body_style, fontName=bold_font, fontSize=12, leading=15),
        ),
    ]
    if supplier.get("kontakt_lice"):
        supplier_lines.append(Paragraph(f"Контакт лице: {supplier['kontakt_lice']}", body_style))
    if supplier.get("adresa"):
        supplier_lines.append(Paragraph(supplier["adresa"], body_style))
    if supplier.get("telefon"):
        supplier_lines.append(Paragraph(f"Телефон: {supplier['telefon']}", body_style))
    if supplier.get("email"):
        supplier_lines.append(Paragraph(f"Email: {supplier['email']}", body_style))

    header_table = Table([[left_header, supplier_lines]], colWidths=[doc.width * 0.5, doc.width * 0.5], hAlign="LEFT")
    header_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#f8fafc")),
                ("BOX", (1, 0), (1, 0), 0.8, colors.HexColor("#cbd5e1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 12),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
            ]
        )
    )
    elements.append(header_table)
    elements.append(Spacer(1, 28))
    elements.append(
        Paragraph(
            "Ве молиме доставете понуда за наведените артикли/предмети подолу. Документот е подготвен од Fersedo како барање за понуда.",
            body_style,
        )
    )
    elements.append(Spacer(1, 24))

    table_rows = [["Р.б.", "Артикл / предмет", "Количина", "Коментар", "Слика"]]
    for index, item in enumerate(items, start=1):
        image_cell = "—"
        image_path = _item_image_filesystem_path(item.get("image"))
        if image_path:
            image_flowable = RLImage(image_path)
            image_flowable._restrictSize(72, 72)
            image_cell = image_flowable
        table_rows.append(
            [
                Paragraph(str(index), table_center_style),
                Paragraph(item.get("item_name") or "—", table_cell_style),
                Paragraph(str(item.get("quantity") or "—"), table_center_style),
                Paragraph(item.get("comment") or "—", table_cell_style),
                image_cell,
            ]
        )

    article_table = Table(
        table_rows,
        colWidths=[34, doc.width * 0.28, 64, doc.width * 0.24, doc.width - 34 - (doc.width * 0.28) - 64 - (doc.width * 0.24)],
        repeatRows=1,
        hAlign="LEFT",
    )
    article_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("FONTNAME", (0, 1), (-1, -1), regular_font),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEADING", (0, 0), (-1, -1), 12),
                ("GRID", (0, 0), (-1, -1), 0.7, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (0, 0), (0, -1), "CENTER"),
                ("ALIGN", (2, 0), (2, -1), "CENTER"),
                ("ALIGN", (4, 0), (4, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(article_table)

    elements.append(Spacer(1, 18))
    footer_table = Table(
        [
            [Paragraph("Подготвил", label_style), Paragraph("Добавувач", label_style)],
            [
                Paragraph(session.get("user", "Fersedo"), body_style),
                Paragraph(supplier.get("naziv") or supplier.get("supplier_name") or "—", body_style),
            ],
        ],
        colWidths=[doc.width / 2, doc.width / 2],
        hAlign="LEFT",
    )
    footer_table.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.HexColor("#cbd5e1")), ("TOPPADDING", (0, 0), (-1, -1), 8)]))
    elements.append(footer_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


@ponudi_bp.route("/dobavuvaci/add", methods=["POST"])
@login_required
def add_supplier():
    if not _is_manager():
        flash("Немате дозвола за додавање добавувачи.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    naziv = request.form.get("naziv", "").strip()
    kontakt_lice = request.form.get("kontakt_lice", "").strip()
    email = request.form.get("email", "").strip()
    telefon = request.form.get("telefon", "").strip()
    adresa = request.form.get("adresa", "").strip()
    zabeleska = request.form.get("zabeleska", "").strip()
    return_draft_id = request.form.get("return_draft_id", "").strip()

    if not naziv:
        flash("Називот на добавувачот е задолжителен.", "warning")
        conn.close()
        return redirect(url_for("ponudi.ponudi", draft_id=return_draft_id or None))

    try:
        cursor.execute(
            """
            INSERT INTO ponudi_dobavuvaci
            (naziv, kontakt_lice, email, telefon, adresa, zabeleska, aktiven)
            VALUES (?, ?, ?, ?, ?, ?, 1)
            """,
            (naziv, kontakt_lice, email, telefon, adresa, zabeleska),
        )
        conn.commit()
        flash(f"Добавувачот {naziv} е успешно креиран.", "success")
    except sqlite3.IntegrityError:
        conn.rollback()
        flash("Добавувач со ова име веќе постои.", "warning")
    except Exception as exc:
        conn.rollback()
        flash(f"Грешка при креирање добавувач: {exc}", "danger")
    finally:
        conn.close()

    return redirect(url_for("ponudi.ponudi", draft_id=return_draft_id or None))


@ponudi_bp.route("/dobavuvaci/delete/<int:supplier_id>", methods=["POST"])
@login_required
def delete_supplier(supplier_id):
    if not _is_manager():
        flash("Немате дозвола за бришење добавувачи.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    linked = cursor.execute(
        "SELECT COUNT(*) AS total FROM ponudi_supplier_links WHERE supplier_id = ?",
        (supplier_id,),
    ).fetchone()
    if linked["total"] > 0:
        conn.close()
        flash("Добавувачот не може да се избрише бидејќи веќе се користи во понуди.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    cursor.execute("DELETE FROM ponudi_dobavuvaci WHERE id = ?", (supplier_id,))
    conn.commit()
    conn.close()
    flash("Добавувачот е избришан.", "success")
    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/dobavuvaci/update/<int:supplier_id>", methods=["POST"])
@login_required
def update_supplier(supplier_id):
    if not _is_manager():
        flash("Немате дозвола за менување добавувачи.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    supplier = cursor.execute("SELECT id FROM ponudi_dobavuvaci WHERE id = ?", (supplier_id,)).fetchone()
    if not supplier:
        conn.close()
        flash("Добавувачот не е пронајден.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    naziv = request.form.get("naziv", "").strip()
    kontakt_lice = request.form.get("kontakt_lice", "").strip()
    email = request.form.get("email", "").strip()
    telefon = request.form.get("telefon", "").strip()
    adresa = request.form.get("adresa", "").strip()
    zabeleska = request.form.get("zabeleska", "").strip()

    if not naziv:
        conn.close()
        flash("Називот на добавувачот е задолжителен.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    duplicate = cursor.execute(
        "SELECT id FROM ponudi_dobavuvaci WHERE lower(naziv) = lower(?) AND id != ?",
        (naziv, supplier_id),
    ).fetchone()
    if duplicate:
        conn.close()
        flash("Веќе постои друг добавувач со ова име.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    try:
        cursor.execute(
            """
            UPDATE ponudi_dobavuvaci
            SET naziv = ?, kontakt_lice = ?, email = ?, telefon = ?, adresa = ?, zabeleska = ?
            WHERE id = ?
            """,
            (naziv, kontakt_lice, email, telefon, adresa, zabeleska, supplier_id),
        )
        cursor.execute(
            """
            UPDATE ponudi_supplier_links
            SET supplier_name = ?, supplier_email = ?
            WHERE supplier_id = ?
            """,
            (naziv, email, supplier_id),
        )
        conn.commit()
        flash(f"Добавувачот {naziv} е успешно ажуриран.", "success")
    except Exception as exc:
        conn.rollback()
        flash(f"Грешка при ажурирање добавувач: {exc}", "danger")
    finally:
        conn.close()

    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/", methods=["GET", "POST"])
@login_required
def ponudi():
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    conn.commit()

    if request.method == "POST":
        action = request.form.get("action", "").strip()
        if action in {"save_draft", "kreiraj"}:
            draft_id = request.form.get("draft_id", type=int)
            draft_row = _load_editable_draft(cursor, draft_id)
            supplier_ids = _parse_supplier_ids(request.form.getlist("supplier_ids"))

            chat_comment = ""
            slika_filename = None

            try:
                items = _parse_item_rows(request)
            except ValueError as exc:
                flash(str(exc), "warning")
                conn.close()
                return redirect(url_for("ponudi.ponudi", draft_id=draft_id or None))

            if items:
                naslov = items[0]["item_name"] if len(items) == 1 else f"Понуда со {len(items)} артикли"
            else:
                naslov = DEFAULT_DRAFT_TITLE

            if action == "kreiraj":
                if not items:
                    flash("Додај барем еден артикл со количина за да ја креираш понудата.", "warning")
                    conn.close()
                    return redirect(url_for("ponudi.ponudi", draft_id=draft_id or None))
                if not supplier_ids:
                    flash("Добавувачот не е креиран или не е избран од листата.", "warning")
                    conn.close()
                    return redirect(url_for("ponudi.ponudi", draft_id=draft_id or None))

            kolicina = sum((item.get("quantity") or 0) for item in items) if items else 0

            try:
                supplier_map = _supplier_map(cursor)
                ponuda_id, ponuda_broj = _save_or_update_ponuda(
                    cursor=cursor,
                    action=action,
                    draft_row=draft_row,
                    slika_filename=slika_filename,
                    title=naslov,
                    kolicina=kolicina,
                    chat_comment=chat_comment,
                    items=items,
                )
                _sync_supplier_links(cursor, ponuda_id, supplier_ids, supplier_map)
                conn.commit()

                if action == "save_draft":
                    flash("Draft понудата е зачувана.", "success")
                    conn.close()
                    return redirect(url_for("ponudi.ponudi", draft_id=ponuda_id))

                now_local = datetime.now().strftime("%d-%m-%Y %H:%M")
                flash(f"Понудата е креирана со број: <strong>{ponuda_broj}</strong> на {now_local}!", "success")

                created_row = cursor.execute("SELECT * FROM ponudi WHERE id = ?", (ponuda_id,)).fetchone()
                ponuda_data = dict(created_row) if created_row else {"id": ponuda_id, "ponuda_broj": ponuda_broj, "username": session["user"], "naslov": naslov}
                ponuda_data["creator_email"] = _user_email(cursor, ponuda_data.get("username"))
                ponuda_data["items"] = _get_items_for_ponuda(cursor, ponuda_id, archived=False, record=ponuda_data)

                sent_count = 0
                failed_suppliers = []
                for supplier_id in supplier_ids:
                    supplier = supplier_map.get(supplier_id)
                    if not supplier:
                        continue
                    send_result = _send_offer_email(ponuda_data, supplier)
                    if send_result["success"]:
                        sent_count += 1
                    else:
                        failed_suppliers.append(f"{supplier.get('naziv') or 'добавувач'} ({send_result['message']})")

                if sent_count and not failed_suppliers:
                    flash(f"Понудата автоматски е испратена до {sent_count} добавувач(и).", "success")
                elif sent_count and failed_suppliers:
                    flash(
                        "Понудата е креирана и делумно испратена. Неуспешно за: "
                        + ", ".join(failed_suppliers),
                        "warning",
                    )
                elif failed_suppliers:
                    flash(
                        "Понудата е креирана, но email не беше испратен: "
                        + ", ".join(failed_suppliers),
                        "warning",
                    )
            except Exception as exc:
                conn.rollback()
                flash(f"Грешка при креирање понуда: {exc}", "danger")
                conn.close()
                return redirect(url_for("ponudi.ponudi", draft_id=draft_id or None))

            conn.close()
            return redirect(url_for("ponudi.ponudi"))

    status_filter = request.args.get("status", "").strip()
    draft_id = request.args.get("draft_id", type=int)

    query = "SELECT * FROM ponudi WHERE COALESCE(is_draft, 0) = 0"
    params = []
    if status_filter:
        query += " AND status = ?"
        params.append(status_filter)
    query += " ORDER BY id DESC"
    rows = cursor.execute(query, params).fetchall()
    ponudi_rows = [_decorate_ponuda(cursor, row) for row in rows]

    drafts_query = "SELECT * FROM ponudi WHERE COALESCE(is_draft, 0) = 1"
    draft_params = []
    if not _is_manager():
        drafts_query += " AND username = ?"
        draft_params.append(session["user"])
    drafts_query += " ORDER BY COALESCE(updated_at, datum_kreiranje) DESC, id DESC"
    draft_rows = cursor.execute(drafts_query, draft_params).fetchall()
    drafts = [_decorate_ponuda(cursor, row, include_comments=False) for row in draft_rows]

    selected_draft = _load_editable_draft(cursor, draft_id)
    selected_draft_data = _decorate_ponuda(cursor, selected_draft) if selected_draft else None

    suppliers = _supplier_rows(cursor)
    supplier_options_json = json.dumps(
        [
            {
                "id": supplier["id"],
                "naziv": supplier["naziv"],
                "kontakt_lice": supplier.get("kontakt_lice") or "",
                "email": supplier.get("email") or "",
                "telefon": supplier.get("telefon") or "",
                "adresa": supplier.get("adresa") or "",
                "zabeleska": supplier.get("zabeleska") or "",
                "aktiven": supplier.get("aktiven", 1),
            }
            for supplier in suppliers
        ],
        ensure_ascii=False,
    )

    conn.close()
    return render_template(
        "ponudi.html",
        ponudi=ponudi_rows,
        drafts=drafts,
        selected_draft=selected_draft_data,
        statusi=STATUSI,
        status_filter=status_filter,
        is_manager=_is_manager(),
        suppliers=suppliers,
        supplier_options_json=supplier_options_json,
        today=date.today().isoformat(),
    )


@ponudi_bp.route("/update_status/<int:ponuda_id>/<string:new_status>")
@login_required
def update_status(ponuda_id, new_status):
    if not _is_manager():
        flash("Немате дозвола!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if new_status not in STATUSI:
        flash("Невалиден статус!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    ponuda = cursor.execute(
        "SELECT * FROM ponudi WHERE id = ? AND COALESCE(is_draft, 0) = 0",
        (ponuda_id,),
    ).fetchone()
    if not ponuda:
        conn.close()
        flash("Понудата не постои!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    try:
        if new_status == "Завршена":
            cursor.execute(
                """
                INSERT OR REPLACE INTO ponudi_archive
                (id, ponuda_broj, username, naslov, dobavuvac, cena, valuta, rok_isporaka,
                 opis, slika, status, datum_kreiranje, datum_vaznost, arhivirano_od,
                 arhivirano_na, kolicina, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, ?, ?)
                """,
                (
                    ponuda["id"],
                    ponuda["ponuda_broj"],
                    ponuda["username"],
                    ponuda["naslov"],
                    ponuda["dobavuvac"],
                    ponuda["cena"],
                    ponuda["valuta"],
                    ponuda["rok_isporaka"],
                    ponuda["opis"],
                    ponuda["slika"],
                    new_status,
                    ponuda["datum_kreiranje"],
                    ponuda["datum_vaznost"],
                    session["user"],
                    ponuda["kolicina"],
                    ponuda["updated_at"],
                ),
            )
            _copy_items_to_archive(cursor, ponuda_id)

            comments = cursor.execute(
                """
                SELECT user, comment, slika, timestamp
                FROM ponudi_comments
                WHERE ponuda_id = ?
                ORDER BY timestamp ASC
                """,
                (ponuda_id,),
            ).fetchall()
            for comment in comments:
                cursor.execute(
                    """
                    INSERT INTO ponudi_archive_comments
                    (archive_ponuda_id, user, comment, slika, timestamp)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (ponuda_id, comment["user"], comment["comment"], comment["slika"], comment["timestamp"]),
                )

            cursor.execute("DELETE FROM ponudi_comments WHERE ponuda_id = ?", (ponuda_id,))
            cursor.execute("DELETE FROM ponudi_supplier_links WHERE ponuda_id = ?", (ponuda_id,))
            cursor.execute("DELETE FROM ponudi_items WHERE ponuda_id = ?", (ponuda_id,))
            cursor.execute("DELETE FROM ponudi WHERE id = ?", (ponuda_id,))
            flash(f"Понудата {ponuda['ponuda_broj']} е архивирана!", "success")
        else:
            cursor.execute(
                "UPDATE ponudi SET status = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                (new_status, ponuda_id),
            )
            flash(f"Статусот е сменет на {new_status}!", "success")

        conn.commit()
    except Exception as exc:
        conn.rollback()
        flash(f"Грешка: {exc}", "danger")
    finally:
        conn.close()

    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/delete_draft/<int:ponuda_id>", methods=["POST"])
@login_required
def delete_draft(ponuda_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    draft = cursor.execute("SELECT * FROM ponudi WHERE id = ? AND COALESCE(is_draft, 0) = 1", (ponuda_id,)).fetchone()
    if not draft:
        conn.close()
        flash("Draft понудата не постои.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    if not (_is_manager() or draft["username"] == session.get("user")):
        conn.close()
        flash("Немате дозвола да ја избришете оваа draft понуда.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    cursor.execute("DELETE FROM ponudi_comments WHERE ponuda_id = ?", (ponuda_id,))
    cursor.execute("DELETE FROM ponudi_supplier_links WHERE ponuda_id = ?", (ponuda_id,))
    cursor.execute("DELETE FROM ponudi_items WHERE ponuda_id = ?", (ponuda_id,))
    cursor.execute("DELETE FROM ponudi WHERE id = ?", (ponuda_id,))
    conn.commit()
    conn.close()
    flash("Draft понудата е избришана.", "success")
    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/add_comment/<int:ponuda_id>", methods=["POST"])
@login_required
def add_comment(ponuda_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    comment = request.form.get("comment", "").strip()
    slika_file = request.files.get("chat_slika")
    ponuda_row = cursor.execute("SELECT username FROM ponudi WHERE id = ?", (ponuda_id,)).fetchone()

    if not _can_comment(ponuda_row):
        conn.close()
        flash("Немате дозвола да коментирате!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if not comment and (not slika_file or not slika_file.filename):
        conn.close()
        flash("Коментарот е празен!", "warning")
        return redirect(url_for("ponudi.ponudi"))

    try:
        chat_slika_filename = None
        if slika_file and slika_file.filename:
            save_dir = os.path.join(current_app.config["STATIC_FOLDER"], "ponudi_chat")
            filename_base = f"chat_{ponuda_id}_{int(time.time())}_{session['user']}"
            chat_slika_filename = _save_compressed_image(slika_file, save_dir, filename_base)

        cursor.execute(
            """
            INSERT INTO ponudi_comments (ponuda_id, user, comment, slika)
            VALUES (?, ?, ?, ?)
            """,
            (ponuda_id, session["user"], comment, chat_slika_filename),
        )
        cursor.execute("UPDATE ponudi SET updated_at = CURRENT_TIMESTAMP WHERE id = ?", (ponuda_id,))
        conn.commit()
        flash("Коментарот е успешно додаден!", "success")
    except Exception as exc:
        conn.rollback()
        flash(f"Грешка: {exc}", "danger")
    finally:
        conn.close()

    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/comments/<int:ponuda_id>")
@login_required
def get_comments(ponuda_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    comments = cursor.execute(
        """
        SELECT user, comment, slika, timestamp
        FROM ponudi_comments
        WHERE ponuda_id = ?
        ORDER BY timestamp ASC
        """,
        (ponuda_id,),
    ).fetchall()
    conn.close()
    return jsonify([dict(comment) for comment in comments])


@ponudi_bp.route("/arhiva")
@login_required
def arhiva():
    if not _is_manager():
        flash("Немате дозвола!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    conn.commit()
    archived_raw = cursor.execute("SELECT * FROM ponudi_archive ORDER BY arhivirano_na DESC").fetchall()

    archived = []
    for row in archived_raw:
        record = dict(row)
        record["datum_kreiranje_display"] = _fmt_datetime(record.get("datum_kreiranje"))
        record["arhivirano_na_display"] = _fmt_datetime(record.get("arhivirano_na"))
        record["comments"] = cursor.execute(
            """
            SELECT user, comment, slika, timestamp
            FROM ponudi_archive_comments
            WHERE archive_ponuda_id = ?
            ORDER BY timestamp ASC
            """,
            (record["id"],),
        ).fetchall()
        archived.append(record)

    conn.close()
    return render_template("ponudi_arhiva.html", archived=archived)


@ponudi_bp.route("/delete_selected", methods=["POST"])
@login_required
def delete_selected():
    if not _is_manager():
        flash("Немате дозвола!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    selected_ids = request.form.getlist("selected_ids")
    if not selected_ids:
        flash("Нема избрани понуди за бришење!", "warning")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    placeholders = ",".join("?" for _ in selected_ids)
    cursor.execute(f"DELETE FROM ponudi_comments WHERE ponuda_id IN ({placeholders})", selected_ids)
    cursor.execute(f"DELETE FROM ponudi_supplier_links WHERE ponuda_id IN ({placeholders})", selected_ids)
    cursor.execute(f"DELETE FROM ponudi_items WHERE ponuda_id IN ({placeholders})", selected_ids)
    cursor.execute(f"DELETE FROM ponudi WHERE id IN ({placeholders})", selected_ids)
    count = cursor.rowcount
    conn.commit()
    conn.close()
    flash(f"Успешно избришани {count} понуди!", "success")
    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/upload_slika/<int:ponuda_id>", methods=["POST"])
@login_required
def upload_slika(ponuda_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    ponuda = cursor.execute("SELECT * FROM ponudi WHERE id = ?", (ponuda_id,)).fetchone()

    if not ponuda:
        conn.close()
        flash("Понудата не постои!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if not (_is_manager() or ponuda["username"] == session.get("user")):
        conn.close()
        flash("Немате дозвола!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if ponuda["slika"]:
        conn.close()
        flash("Оваа понуда веќе има прикачена слика!", "warning")
        return redirect(url_for("ponudi.ponudi"))

    slika_file = request.files.get("nova_slika")
    if not slika_file or not slika_file.filename:
        conn.close()
        flash("Нема избрана слика!", "warning")
        return redirect(url_for("ponudi.ponudi"))

    try:
        save_dir = os.path.join(current_app.config["STATIC_FOLDER"], "ponudi")
        filename_base = f"pon_{ponuda_id}_{int(time.time())}_{session['user']}"
        slika_filename = _save_compressed_image(slika_file, save_dir, filename_base)
        if not slika_filename:
            conn.close()
            flash("Грешка при зачувување на сликата!", "danger")
            return redirect(url_for("ponudi.ponudi"))

        cursor.execute(
            "UPDATE ponudi SET slika = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (slika_filename, ponuda_id),
        )
        conn.commit()
        flash(f"Сликата е успешно прикачена на понуда {ponuda['ponuda_broj'] or ponuda['naslov']}!", "success")
    except Exception as exc:
        conn.rollback()
        flash(f"Грешка: {exc}", "danger")
    finally:
        conn.close()

    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/pdf/<int:ponuda_id>/<int:supplier_id>")
@login_required
def offer_pdf(ponuda_id, supplier_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    ponuda = cursor.execute(
        "SELECT * FROM ponudi WHERE id = ? AND COALESCE(is_draft, 0) = 0",
        (ponuda_id,),
    ).fetchone()
    if not ponuda:
        conn.close()
        flash("Понудата не постои или е во draft.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if not (_is_manager() or ponuda["username"] == session.get("user")):
        conn.close()
        flash("Немате дозвола за оваа PDF понуда.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    link = cursor.execute(
        """
        SELECT supplier_id, supplier_name, supplier_email
        FROM ponudi_supplier_links
        WHERE ponuda_id = ? AND supplier_id = ?
        """,
        (ponuda_id, supplier_id),
    ).fetchone()
    if not link:
        conn.close()
        flash("Избраниот добавувач не е поврзан со оваа понуда.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    supplier = cursor.execute("SELECT * FROM ponudi_dobavuvaci WHERE id = ?", (supplier_id,)).fetchone()
    supplier_data = dict(supplier) if supplier else dict(link)
    ponuda_data = dict(ponuda)
    ponuda_data["creator_email"] = _user_email(cursor, ponuda["username"])
    ponuda_data["items"] = _get_items_for_ponuda(cursor, ponuda_id, archived=False, record=ponuda_data)
    conn.close()

    pdf_buffer = _generate_offer_pdf_v2(ponuda_data, supplier_data)
    supplier_name = supplier_data.get("naziv") or supplier_data.get("supplier_name") or "dobavuvac"
    filename = f"{_safe_filename_part(ponuda['ponuda_broj'] or ponuda['naslov'])}_{_safe_filename_part(supplier_name)}.pdf"
    return send_file(pdf_buffer, mimetype="application/pdf", as_attachment=True, download_name=filename)


@ponudi_bp.route("/send/<int:ponuda_id>/<int:supplier_id>", methods=["POST"])
@login_required
def send_offer_email(ponuda_id, supplier_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    ponuda = cursor.execute(
        "SELECT * FROM ponudi WHERE id = ? AND COALESCE(is_draft, 0) = 0",
        (ponuda_id,),
    ).fetchone()
    if not ponuda:
        conn.close()
        flash("Понудата не постои или е во draft.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if not (_is_manager() or ponuda["username"] == session.get("user")):
        conn.close()
        flash("Немате дозвола за испраќање на оваа понуда.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    link = cursor.execute(
        """
        SELECT supplier_id, supplier_name, supplier_email
        FROM ponudi_supplier_links
        WHERE ponuda_id = ? AND supplier_id = ?
        """,
        (ponuda_id, supplier_id),
    ).fetchone()
    if not link:
        conn.close()
        flash("Избраниот добавувач не е поврзан со оваа понуда.", "warning")
        return redirect(url_for("ponudi.ponudi"))

    supplier = cursor.execute("SELECT * FROM ponudi_dobavuvaci WHERE id = ?", (supplier_id,)).fetchone()
    supplier_data = dict(supplier) if supplier else dict(link)
    ponuda_data = dict(ponuda)
    ponuda_data["creator_email"] = _user_email(cursor, ponuda["username"])
    ponuda_data["items"] = _get_items_for_ponuda(cursor, ponuda_id, archived=False, record=ponuda_data)
    conn.close()

    result = _send_offer_email(ponuda_data, supplier_data)
    flash(result["message"], "success" if result["success"] else "danger")
    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/send_all/<int:ponuda_id>", methods=["POST"])
@login_required
def send_offer_email_all(ponuda_id):
    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)

    ponuda = cursor.execute(
        "SELECT * FROM ponudi WHERE id = ? AND COALESCE(is_draft, 0) = 0",
        (ponuda_id,),
    ).fetchone()
    if not ponuda:
        conn.close()
        flash("Понудата не постои или е во draft.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    if not (_is_manager() or ponuda["username"] == session.get("user")):
        conn.close()
        flash("Немате дозвола за испраќање на оваа понуда.", "danger")
        return redirect(url_for("ponudi.ponudi"))

    links = _get_links_for_ponuda(cursor, ponuda_id)
    supplier_map = _supplier_map(cursor)
    ponuda_data = dict(ponuda)
    ponuda_data["creator_email"] = _user_email(cursor, ponuda["username"])
    ponuda_data["items"] = _get_items_for_ponuda(cursor, ponuda_id, archived=False, record=ponuda_data)
    conn.close()

    sent = 0
    failed = []
    for link in links:
        supplier = supplier_map.get(link["supplier_id"], link)
        result = _send_offer_email(ponuda_data, supplier)
        if result["success"]:
            sent += 1
        else:
            failed.append((supplier.get("naziv") or supplier.get("supplier_name") or "добавувач", result["message"]))

    if sent and not failed:
        flash(f"Понудата е испратена до сите избрани добавувачи ({sent}).", "success")
    elif sent and failed:
        failed_names = ", ".join(name for name, _ in failed)
        flash(f"Испратено до {sent} добавувачи, но неуспешно за: {failed_names}.", "warning")
    else:
        flash("Понудата не беше испратена до ниеден добавувач.", "danger")
    return redirect(url_for("ponudi.ponudi"))


@ponudi_bp.route("/export/excel")
@login_required
def export_excel():
    if not _is_manager():
        flash("Немате дозвола!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    rows = cursor.execute(
        "SELECT * FROM ponudi WHERE COALESCE(is_draft, 0) = 0 ORDER BY id DESC"
    ).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Понуди"
    ws.append(["Број", "Корисник", "Наслов", "Добавувачи", "Количина", "Статус", "Креирана", "Последна промена"])
    for row in rows:
        ws.append(
            [
                row["ponuda_broj"],
                row["username"],
                row["naslov"],
                row["dobavuvac"],
                row["kolicina"],
                row["status"],
                row["datum_kreiranje"],
                row["updated_at"],
            ]
        )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"ponudi_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
    )


@ponudi_bp.route("/export/pdf")
@login_required
def export_pdf():
    if not _is_manager():
        flash("Немате дозвола!", "danger")
        return redirect(url_for("ponudi.ponudi"))

    conn = get_db()
    cursor = conn.cursor()
    _ensure_tables(cursor)
    rows = cursor.execute(
        "SELECT * FROM ponudi WHERE COALESCE(is_draft, 0) = 0 ORDER BY id DESC"
    ).fetchall()
    conn.close()

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    data = [["Број", "Наслов", "Добавувачи", "Кол.", "Статус"]]
    for row in rows:
        data.append(
            [
                row["ponuda_broj"],
                row["naslov"],
                row["dobavuvac"] or "—",
                row["kolicina"] or "—",
                row["status"],
            ]
        )

    table = Table(data, colWidths=[50, 120, 190, 45, 70])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ]
        )
    )

    doc.build([Paragraph("Листа на понуди", styles["Heading1"]), Spacer(1, 12), table])
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"ponudi_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
    )
